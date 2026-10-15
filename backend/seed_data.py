import os
import sys
from pathlib import Path
import openpyxl
from backend.database import get_db_connection, init_db

if sys.platform == "win32":
    try:
        sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    except Exception:
        pass

DATA_FILE_PATH = Path(__file__).resolve().parent.parent / "data" / "events_seed.xlsx"


def seed_events() -> int:
    """
    Seeds the events table from data/events_seed.xlsx.
    Skips seeding if the events table already has rows.
    Raises FileNotFoundError if data/events_seed.xlsx does not exist.
    Returns the number of events inserted (or current count if already seeded).
    """
    init_db()

    conn = get_db_connection()
    cursor = conn.cursor()

    # Check if events table already contains data
    cursor.execute("SELECT COUNT(*) AS count FROM events")
    current_count = cursor.fetchone()["count"]

    if current_count > 0:
        print(f"Events table already has {current_count} rows. Skipping seed.")
        conn.close()
        return current_count

    # Verify seed file exists
    if not DATA_FILE_PATH.exists():
        conn.close()
        raise FileNotFoundError(
            f"Seed data file not found at: {DATA_FILE_PATH}\n"
            "Please copy the provided events_seed.xlsx file into the data/ folder before running this script."
        )

    wb = openpyxl.load_workbook(str(DATA_FILE_PATH))
    if "Events" not in wb.sheetnames:
        conn.close()
        raise ValueError("Spreadsheet does not contain a sheet named 'Events'.")

    sheet = wb["Events"]
    rows = list(sheet.iter_rows(values_only=True))

    if not rows or len(rows) < 2:
        conn.close()
        raise ValueError("Spreadsheet 'Events' sheet is empty or contains no data rows.")

    header = rows[0]
    expected_headers = [
        "id", "title", "category", "description",
        "event_date", "event_time", "venue", "capacity", "fee", "icon"
    ]

    header_clean = [str(h).strip().lower() if h is not None else "" for h in header]
    if header_clean[:len(expected_headers)] != expected_headers:
        print(f"Warning: Header row is {header_clean}, expecting {expected_headers}")

    inserted = 0
    for row in rows[1:]:
        if not row or row[0] is None:
            continue
        (
            event_id,
            title,
            category,
            description,
            event_date,
            event_time,
            venue,
            capacity,
            fee,
            icon,
        ) = row[:10]

        cursor.execute(
            """
            INSERT INTO events (id, title, category, description, event_date, event_time, venue, capacity, fee, icon)
            VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
            """,
            (
                int(event_id),
                str(title),
                str(category),
                str(description) if description is not None else "",
                str(event_date),
                str(event_time),
                str(venue),
                int(capacity),
                int(fee) if fee is not None else 0,
                str(icon) if icon is not None else "",
            ),
        )
        inserted += 1

    conn.commit()
    conn.close()
    print(f"Successfully seeded {inserted} events into the database.")
    return inserted


if __name__ == "__main__":
    count = seed_events()
    conn = get_db_connection()
    cursor = conn.cursor()
    cursor.execute("SELECT id, title, category, event_date, event_time, venue, capacity, fee, icon FROM events ORDER BY id ASC")
    rows = cursor.fetchall()
    print(f"\n--- Events in fest.db ({len(rows)} rows) ---")
    for r in rows:
        print(f"[{r['id']}] {r['icon']} {r['title']} | {r['category']} | {r['event_date']} {r['event_time']} | {r['venue']} | Cap: {r['capacity']} | Fee: ₹{r['fee']}")
    conn.close()
